from docx import Document


def extract_text_from_docx(file_path):
    """
    Extract text from a DOCX file.

    Handles:
    1. Normal paragraphs
    2. Tables

    Returns:
        Complete extracted text as a single string.
    """

    document = Document(file_path)

    extracted_parts = []

    # ---------------------------------------------------------
    # 1. Extract normal paragraphs
    # ---------------------------------------------------------

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:
            extracted_parts.append(text)

    # ---------------------------------------------------------
    # 2. Extract tables
    # ---------------------------------------------------------

    for table in document.tables:

        for row in table.rows:

            row_values = []

            for cell in row.cells:

                cell_text = cell.text.strip()

                if cell_text:
                    row_values.append(cell_text)

            if row_values:

                extracted_parts.append(
                    " | ".join(row_values)
                )

    # ---------------------------------------------------------
    # 3. Combine everything
    # ---------------------------------------------------------

    return "\n".join(
        extracted_parts
    )